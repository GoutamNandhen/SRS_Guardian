from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
import docx
import pymupdf


@dataclass
class ExtractedDocument:
    text: str
    page_count: int
    meta: dict[str, Any] = field(default_factory=dict)
    pages: list[str] = field(default_factory=list)


def extract_pdf(path: str | Path) -> ExtractedDocument:
    file_path = Path(path)
    doc = pymupdf.open(file_path)
    pages: list[str] = []

    for page_idx in range(len(doc)):
        page = doc.load_page(page_idx)
        page_text = page.get_text("text")
        pages.append(page_text)

    doc.close()
    full_text = "\n\n".join(pages)

    return ExtractedDocument(
        text=full_text,
        page_count=len(pages),
        meta={
            "filename": file_path.name,
            "file_type": "pdf",
            "file_size": file_path.stat().st_size,
        },
        pages=pages,
    )


def extract_docx(path: str | Path) -> ExtractedDocument:
    file_path = Path(path)
    doc = docx.Document(file_path)
    lines: list[str] = []

    for paragraph in doc.paragraphs:
        cleaned_p = paragraph.text.strip()
        if cleaned_p:
            lines.append(cleaned_p)

    for table in doc.tables:
        for row in table.rows:
            seen_cells: list[str] = []
            for cell in row.cells:
                cell_content = cell.text.strip()
                if cell_content and (not seen_cells or cell_content != seen_cells[-1]):
                    seen_cells.append(cell_content)
            if seen_cells:
                lines.append(" | ".join(seen_cells))

    full_text = "\n\n".join(lines)
    pages = [full_text] if full_text else []

    return ExtractedDocument(
        text=full_text,
        page_count=max(1, len(pages)),
        meta={
            "filename": file_path.name,
            "file_type": "docx",
            "file_size": file_path.stat().st_size,
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        },
        pages=pages,
    )


def extract_txt(path: str | Path) -> ExtractedDocument:
    file_path = Path(path)
    try:
        content = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = file_path.read_text(encoding="latin-1", errors="replace")

    return ExtractedDocument(
        text=content,
        page_count=1,
        meta={
            "filename": file_path.name,
            "file_type": "txt",
            "file_size": file_path.stat().st_size,
        },
        pages=[content] if content else [],
    )


def extract_text(path: str | Path) -> ExtractedDocument:
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf(file_path)
    elif suffix == ".docx":
        return extract_docx(file_path)
    elif suffix == ".txt":
        return extract_txt(file_path)
    else:
        raise ValueError(
            f"Unsupported file format '{suffix}'. Supported formats: .pdf, .docx, .txt"
        )
